"""Assemble manuscript: markdown sections -> citation resolution -> .docx (+ combined .md, word count)."""
import re, json, sys, os, glob
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
M = "manuscript/"
refs = {r["key"]: r for r in json.load(open(M + "refs_verified.json"))}
order = ["front.md", "sec1_intro.md", "sec2_related.md", "sec3_data.md", "sec4_method.md", "sec5_protocol.md", "sec6a.md", "sec6b.md", "sec6c.md", "sec6d.md", "sec7_discussion.md", "sec8_conclusion.md", "back.md"]
text = "\n\n".join(open(M + f).read() for f in order if os.path.exists(M + f))
text = re.sub(r"@@TBL:(\w+)@@", lambda m: open(M + f"tbl_{m.group(1)}.md").read(), text)
figs = re.findall(r"!\[[^\]]*\]\((figures/fig\d+_(\w+)\.png)\)", text); figno = {k: i + 1 for i, (_, k) in enumerate(figs)}
text = re.sub(r"@@fig:(\w+)@@", lambda m: f"Fig. {figno[m.group(1)]}", text)
text = re.sub(r"Fig\. @@num:(\w+)@@", lambda m: f"Fig. {figno[m.group(1)]}", text)
_n = [0]
def _cap(m):
    _n[0] += 1; return f"![Figure {_n[0]}. {m.group(1)}]({m.group(2)})"
text = re.sub(r"!\[([^\]]*)\]\((figures/[^)]+)\)", _cap, text)
print("figure order:", [k for _, k in figs])
cited = []
def surname(a): return a.split(",")[0].strip()
def short(k):
    r = refs[k]; au = r["authors"]
    if len(au) == 1: n = surname(au[0])
    elif len(au) == 2: n = f"{surname(au[0])} and {surname(au[1])}"
    else: n = f"{surname(au[0])} et al."
    return f"{n}, {r['year']}"
def cite(m):
    keys = [k.strip().lstrip("@") for k in m.group(1).split(";")]
    for k in keys:
        if k not in refs: print("MISSING REF", k); return "(??)"
        if k not in cited: cited.append(k)
    keys.sort(key=lambda k: (short(k)))
    return "(" + "; ".join(short(k) for k in keys) + ")"
text = re.sub(r"\[((?:@[\w]+\s*;?\s*)+)\]", cite, text)
def fmt_ref(r):
    au = r["authors"]; a = ", ".join(au[:-1]) + (" & " if len(au) > 1 else "") + au[-1] if len(au) <= 6 else ", ".join(au[:6]) + ", et al."
    s = f"{a} ({r['year']}). {r['title']}. {r['venue']}"
    if r.get("volume"): s += f", {r['volume']}" + (f"({r['issue']})" if r.get("issue") else "")
    if r.get("pages"): s += f", {r['pages']}"
    s += "."
    if r.get("doi"): s += f" https://doi.org/{r['doi']}"
    return s
bib = sorted(cited, key=lambda k: refs[k]["authors"][0].lower() + str(refs[k]["year"]))
text += "\n\n# References\n\n" + "\n\n".join("@@REF@@" + fmt_ref(refs[k]) for k in bib)
open(M + "manuscript_full.md", "w").write(text.replace("@@REF@@", ""))
body = re.sub(r"```.*?```", "", text, flags=re.S)
words_all = len(re.findall(r"\b[\w\-’'.,%]+\b", body))
main = body.split("# References")[0]
words_main = len(re.findall(r"[A-Za-z0-9][\w\-’'.,%/]*", main))
refw = len(re.findall(r"[A-Za-z0-9][\w\-’'.,%/]*", body.split("# References")[1]))
print(f"cited refs {len(cited)}; words (main text incl tables, excl references): {words_main}; references: {refw}")
# ---------------- docx
doc = Document()
sec = doc.sections[0]; sec.left_margin = sec.right_margin = Cm(2.2); sec.top_margin = sec.bottom_margin = Cm(2.2)
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.15
for lvl, sz in ((1, 15), (2, 12.5), (3, 11.5)):
    h = doc.styles[f"Heading {lvl}"]; h.font.name = "Times New Roman"; h.font.size = Pt(sz); h.font.bold = True; h.font.color.rgb = RGBColor(0, 0, 0)
def runs(p, s, size=None):
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", s):
        if not part: continue
        if part.startswith("**"): r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*") and len(part) > 2: r = p.add_run(part[1:-1]); r.italic = True
        else: r = p.add_run(part)
        if size: r.font.size = Pt(size)
def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color); tcPr.append(shd)
lines = text.split("\n"); i = 0
while i < len(lines):
    ln = lines[i]
    if ln.startswith("@@REF@@"):
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.first_line_indent = Cm(-0.8); p.paragraph_format.space_after = Pt(3)
        runs(p, ln[7:], 9.5)
    elif ln.startswith("TITLE: "):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run(ln[7:]); r.bold = True; r.font.size = Pt(17)
    elif ln.startswith("CENTER: "):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; runs(p, ln[8:])
    elif ln.startswith("### "): doc.add_heading(ln[4:], 3)
    elif ln.startswith("## "): doc.add_heading(ln[3:], 2)
    elif ln.startswith("# "): doc.add_heading(ln[2:], 1)
    elif ln.startswith("PAGEBREAK"): doc.add_page_break()
    elif ln.startswith("!["):
        m = re.match(r"!\[(.*?)\]\((.*?)\)", ln)
        if m and os.path.exists(m.group(2)):
            doc.add_picture(m.group(2), width=Inches(6.2)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; runs(p, m.group(1), 9.5)
        else: print("missing figure", ln[:80])
    elif ln.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|---"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            if not lines[i].startswith("|---"): rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
            i += 1
        nc = max(len(r) for r in rows); t = doc.add_table(rows=len(rows), cols=nc); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, r in enumerate(rows):
            for ci in range(nc):
                c = t.cell(ri, ci); c.text = ""; pp = c.paragraphs[0]; pp.paragraph_format.space_after = Pt(0)
                runs(pp, r[ci] if ci < len(r) else "", 8.5)
                if ri == 0:
                    shade(c, "E6E5E1")
                    for rr in pp.runs: rr.bold = True
        doc.add_paragraph(); continue
    elif re.match(r"^\d+\. ", ln): p = doc.add_paragraph(style="List Number"); runs(p, re.sub(r"^\d+\. ", "", ln))
    elif ln.startswith("- "): p = doc.add_paragraph(style="List Bullet"); runs(p, ln[2:])
    elif ln.strip() == "": pass
    else:
        p = doc.add_paragraph(); runs(p, ln)
    i += 1
doc.save(M + "manuscript.docx"); print("saved")
